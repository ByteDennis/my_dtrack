from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from . import companion_table, validate_table_name

PAIR_WORKBOOK_COLUMNS = (
    "pair_name",
    "enabled",
    "config_json",
    "left_sql",
    "right_sql",
    "vintage_sample_size",
    "vintage_values",
)


def create_pair_workbook(path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "pairs"
    sheet.append(PAIR_WORKBOOK_COLUMNS)
    sheet.append([None] * len(PAIR_WORKBOOK_COLUMNS))
    sheet.freeze_panes = "A2"
    descriptions = {
        "pair_name": "Unique key shared by pair metadata and pair SQL.",
        "enabled": "TRUE or FALSE. Only enabled pairs appear in the run prompt.",
        "config_json": "JSON object containing Left, Right, key, date, sample, and check settings.",
        "left_sql": "Optional Oracle/Hadoop base query. Keep {{query_windows:name}} markers intact.",
        "right_sql": "Optional Databricks base query. Keep {{query_windows:name}} markers intact.",
        "vintage_sample_size": "Optional positive integer. Leave blank when vintage_values is used.",
        "vintage_values": 'Optional JSON array such as ["2026-01-01","2026-04-01"].',
    }
    widths = (24, 12, 90, 70, 70, 24, 42)
    for index, (name, width) in enumerate(zip(PAIR_WORKBOOK_COLUMNS, widths), 1):
        cell = sheet.cell(1, index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.comment = Comment(descriptions[name], "PairRecon")
        sheet.column_dimensions[cell.column_letter].width = width
    table = Table(displayName="PairReconPairs", ref="A1:G2")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    temporary = target.with_name(f".{target.name}.part")
    workbook.save(temporary)
    temporary.replace(target)
    return target


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _boolean(value: Any, *, row: int) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, int | float) and value in {0, 1}:
        return bool(value)
    text = str(value or "").strip().casefold()
    if text in {"true", "yes", "y", "1"}:
        return True
    if text in {"false", "no", "n", "0"}:
        return False
    raise ValueError(f"pairs.xlsx row {row}: enabled must be true or false")


def _config(value: Any, *, row: int) -> str:
    try:
        parsed = json.loads(str(value))
    except (TypeError, json.JSONDecodeError) as exc:
        raise ValueError(f"pairs.xlsx row {row}: config_json must be valid JSON") from exc
    if not isinstance(parsed, dict):
        raise TypeError(f"pairs.xlsx row {row}: config_json must be a JSON object")
    return json.dumps(parsed, ensure_ascii=False, separators=(",", ":"), sort_keys=True)


def _sample_size(value: Any, *, row: int) -> int | None:
    if _blank(value):
        return None
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise ValueError(f"pairs.xlsx row {row}: vintage_sample_size must be an integer") from None
    if number <= 0 or (isinstance(value, float) and not value.is_integer()):
        raise ValueError(f"pairs.xlsx row {row}: vintage_sample_size must be positive")
    return number


def _vintage_values(value: Any, *, row: int) -> list[str] | None:
    if _blank(value):
        return None
    try:
        parsed = json.loads(str(value))
    except json.JSONDecodeError as exc:
        raise ValueError(f"pairs.xlsx row {row}: vintage_values must be a JSON array") from exc
    if (
        not isinstance(parsed, list)
        or not parsed
        or not all(isinstance(item, str) and item.strip() for item in parsed)
    ):
        raise ValueError(
            f"pairs.xlsx row {row}: vintage_values must be a non-empty array of strings"
        )
    return [item.strip() for item in parsed]


def read_pair_workbook(path: str | Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    source = Path(path)
    workbook = load_workbook(source, read_only=False, data_only=True)
    if workbook.sheetnames != ["pairs"]:
        raise ValueError("pairs.xlsx must contain exactly one worksheet named 'pairs'")
    sheet = workbook["pairs"]
    rows = sheet.iter_rows(values_only=True)
    header = tuple(str(value or "").strip() for value in next(rows, ()))
    if header != PAIR_WORKBOOK_COLUMNS:
        raise ValueError("pairs.xlsx columns must be exactly: " + ", ".join(PAIR_WORKBOOK_COLUMNS))
    result = []
    seen = set()
    for number, values in enumerate(rows, 2):
        if all(_blank(value) for value in values):
            continue
        row = dict(zip(PAIR_WORKBOOK_COLUMNS, values))
        pair = str(row["pair_name"] or "").strip()
        if not pair:
            raise ValueError(f"pairs.xlsx row {number}: pair_name is required")
        if pair in seen:
            raise ValueError(f"pairs.xlsx row {number}: duplicate pair_name {pair!r}")
        seen.add(pair)
        sample_size = _sample_size(row["vintage_sample_size"], row=number)
        vintage_values = _vintage_values(row["vintage_values"], row=number)
        if sample_size is not None and vintage_values is not None:
            raise ValueError(
                f"pairs.xlsx row {number}: set vintage_sample_size or vintage_values, not both"
            )
        result.append(
            {
                "pair_name": pair,
                "enabled": _boolean(row["enabled"], row=number),
                "config_json": _config(row["config_json"], row=number),
                "left_sql": None if _blank(row["left_sql"]) else str(row["left_sql"]).strip(),
                "right_sql": (None if _blank(row["right_sql"]) else str(row["right_sql"]).strip()),
                "vintage_sample_size": sample_size,
                "vintage_values": vintage_values,
            }
        )
    return result


def upsert_pair_configuration(sql, pair_table: str, rows: list[dict[str, Any]]) -> dict:
    pair_table = validate_table_name(pair_table)
    pair_sql_table = companion_table(pair_table, "pair_recon_pair_sql")
    meta_statement = f"""MERGE INTO {pair_table} AS t
USING (SELECT :pair_name AS pair_name) AS s ON t.pair_name = s.pair_name
WHEN MATCHED THEN UPDATE SET
  t.enabled = CAST(:enabled AS BOOLEAN),
  t.config_json = :config_json,
  t.updated_at = CAST(current_timestamp() AS TIMESTAMP_NTZ)
WHEN NOT MATCHED THEN INSERT (
  pair_name, enabled, config_json, latest_run_id, latest_status, updated_at
) VALUES (
  :pair_name, CAST(:enabled AS BOOLEAN), :config_json, NULL, NULL,
  CAST(current_timestamp() AS TIMESTAMP_NTZ)
)"""
    pair_sql_statement = f"""MERGE INTO {pair_sql_table} AS t
USING (SELECT :pair_name AS pair_name) AS s ON t.pair_name = s.pair_name
WHEN MATCHED THEN UPDATE SET
  t.left_sql = :left_sql,
  t.right_sql = :right_sql,
  t.vintage_sample_size = CAST(:vintage_sample_size AS INT),
  t.vintage_values = from_json(:vintage_values, 'ARRAY<STRING>')
WHEN NOT MATCHED THEN INSERT (
  pair_name, left_sql, right_sql, vintage_sample_size, vintage_values
) VALUES (
  :pair_name, :left_sql, :right_sql, CAST(:vintage_sample_size AS INT),
  from_json(:vintage_values, 'ARRAY<STRING>')
)"""
    for row in rows:
        values = {
            **row,
            "enabled": "true" if row["enabled"] else "false",
            "vintage_values": (
                json.dumps(row["vintage_values"], ensure_ascii=False)
                if row["vintage_values"] is not None
                else None
            ),
        }
        sql.execute(meta_statement, values)
        sql.execute(pair_sql_statement, values)
    return {
        "pair_meta_table": pair_table,
        "pair_sql_table": pair_sql_table,
        "rows": len(rows),
        "pairs": [row["pair_name"] for row in rows],
    }
