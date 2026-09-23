from __future__ import annotations

import re
from typing import Any

_ILLEGAL = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")
STAT_ROWS = (
    ("Type", "col_type"),
    ("N_Total", "n_total"),
    ("N_Missing", "n_missing"),
    ("N_Unique", "n_unique"),
    ("Mean", "mean"),
    ("Std", "std"),
    ("Min", "min"),
    ("Max", "max"),
    ("Length avg", "length_avg"),
    ("Length max", "length_max"),
    ("Hash sum", "hash_sum"),
)


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        return _ILLEGAL.sub("", value)
    return value


def _number(value: Any, field: str) -> Any:
    if value in (None, ""):
        return None
    if field == "hash_sum":
        return _clean(str(value))
    try:
        return float(value)
    except (TypeError, ValueError):
        return _clean(value)


def _differs(row: dict[str, Any], field: str) -> bool:
    if field == "col_type":
        return False
    if field in ("n_total", "n_missing", "n_unique"):
        return bool(row.get(f"{field}_diff"))
    return row.get(f"{field}_match") is False


# >>> Stage 1 sheet: header block then one line per Partition, findings first and in red <<< #
def write_row_sheet(workbook, document: dict[str, Any]) -> None:
    from openpyxl.styles import Font, PatternFill

    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    red_font = Font(color="CC0000")
    bold = Font(bold=True)
    left, right = document["left"], document["right"]
    left_label = str(left.get("source") or "left").upper()
    right_label = str(right.get("source") or "databricks").upper()
    summary = document.get("summary") or {}
    sheet = workbook.create_sheet(title=f"{document['pair']}_row"[:31])
    header = [
        (f"{left_label}:", _relation(left)),
        (f"{right_label}:", _relation(right)),
        ("Date Variable:", f"{left.get('date_col')} / {right.get('date_col')}"),
        (
            "Bounds:",
            f"{(document.get('bounds') or {}).get('from')} -> {(document.get('bounds') or {}).get('to')}",
        ),
        ("Matching:", summary.get("match", 0)),
        ("Mismatch:", summary.get("count_diff", 0)),
        (f"{left_label} Only:", summary.get("left_only", 0)),
        (f"{right_label} Only:", summary.get("right_only", 0)),
        ("Gate:", str(document.get("gate_decision"))),
    ]
    for index, (label, value) in enumerate(header, start=1):
        sheet.cell(row=index, column=1, value=label).font = bold
        sheet.cell(row=index, column=2, value=_clean(value))
    row_num = len(header) + 2
    for column, label in enumerate(("Date", left_label, right_label, "Diff", "Status"), start=1):
        sheet.cell(row=row_num, column=column, value=label).font = bold
    row_num += 1
    order = {"count_diff": 0, "left_only": 1, "right_only": 2, "match": 3}
    partitions = sorted(
        document["partitions"].items(), key=lambda item: (order[item[1]["verdict"]], item[0])
    )
    for partition, row in partitions:
        values = (partition, row.get("left"), row.get("right"), row.get("delta"), row["verdict"])
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_num, column=column, value=_clean(value))
            if row["verdict"] != "match":
                cell.fill = red_fill
                cell.font = red_font
        row_num += 1


def _relation(side: dict[str, Any]) -> str:
    parts = [side.get(name) for name in ("catalog", "schema", "table")]
    return ".".join(str(part) for part in parts if part) or str(side.get("path") or "")


# >>> Stage 2 sheet: per Window, both sides transposed with diff cells in red <<< #
def write_col_sheet(
    workbook, pair: str, rows: list[dict[str, Any]], left_label: str, right_label: str
) -> None:
    from openpyxl.styles import Font, PatternFill

    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    bold = Font(bold=True)
    sheet = workbook.create_sheet(title=f"{pair}_col"[:31])
    row_num = 1
    by_vintage: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        by_vintage.setdefault(row["vintage"], []).append(row)
    for vintage in sorted(by_vintage):
        comps = sorted(
            by_vintage[vintage], key=lambda row: (row.get("verdict") != "diff", row["column"])
        )
        active = [
            (label, field)
            for label, field in STAT_ROWS
            if any(
                comp.get(f"{field}_left") not in (None, "")
                or comp.get(f"{field}_right") not in (None, "")
                or field == "col_type"
                for comp in comps
            )
        ]
        has_diff = any(comp.get("verdict") == "diff" for comp in comps)
        sheet.cell(row=row_num, column=1, value=f"Vintage: {vintage}").font = bold
        status = sheet.cell(row=row_num, column=2, value="FAIL" if has_diff else "PASS")
        status.fill = red_fill if has_diff else green_fill
        status.font = bold
        row_num += 1
        for side, label, name_key in (
            ("left", left_label.upper(), "column"),
            ("right", right_label.upper(), "right_col"),
        ):
            sheet.cell(row=row_num, column=1, value=label).font = bold
            row_num += 1
            for index, comp in enumerate(comps):
                name = comp.get(name_key) or comp["column"]
                sheet.cell(row=row_num, column=index + 2, value=_clean(name))
            row_num += 1
            for stat_label, field in active:
                sheet.cell(row=row_num, column=1, value=stat_label).font = bold
                for index, comp in enumerate(comps):
                    value = (
                        comp.get("col_type") if field == "col_type" else comp.get(f"{field}_{side}")
                    )
                    cell = sheet.cell(row=row_num, column=index + 2, value=_number(value, field))
                    if _differs(comp, field):
                        cell.fill = red_fill
                    elif comp.get("verdict") == "diff" and field != "col_type":
                        cell.fill = green_fill
                row_num += 1
            row_num += 1
        row_num += 1


# >>> one workbook: <pair>_row and <pair>_col sheets for every Pair <<< #
def build_workbook(
    documents: list[dict[str, Any]], results_by_pair: dict[str, list[dict[str, Any]]]
):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for document in documents:
        write_row_sheet(workbook, document)
        rows = results_by_pair.get(document["pair"]) or []
        if rows:
            write_col_sheet(
                workbook,
                document["pair"],
                rows,
                str(document["left"].get("source") or "left"),
                str(document["right"].get("source") or "databricks"),
            )
    return workbook
