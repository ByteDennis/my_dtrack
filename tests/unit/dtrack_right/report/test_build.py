import json

import pytest
from support import MemoryStore

from dtrack_left.backends.store import fetch_report
from dtrack_right.backends.store import VolumeStore
from dtrack_right.checks.col_check import compare_statistics
from dtrack_right.report import build_report
from dtrack_right.report.build import compare_col_csv_bytes

LEFT = {
    "col_type": "numeric",
    "n_total": 3,
    "n_missing": 0,
    "n_unique": 3,
    "hash_sum": "10",
    "mean": 2.0,
    "std": 1.0,
    "min_val": "1",
    "max_val": "3",
    "length_avg": None,
    "length_max": None,
    "top_10": "",
}


def _store():
    store = MemoryStore()
    store.put_json("row_check/p.left.json", {"pair": "p", "counts": {}})
    store.put_json(
        "row_check/p.json",
        {
            "run_id": "run",
            "pair": "p",
            "left": {"source": "oracle", "table": "A", "date_col": "DT"},
            "right": {
                "source": "databricks",
                "catalog": "c",
                "schema": "s",
                "table": "a",
                "date_col": "dt",
            },
            "bounds": {"from": "2026-01-01", "to": "2026-01-03"},
            "partitions": {
                "2026-01-01": {"left": 1, "right": 1, "verdict": "match"},
                "2026-01-02": {"left": 1, "right": 2, "verdict": "count_diff", "delta": 1},
                "2026-01-03": {"left": 1, "right": None, "verdict": "left_only"},
            },
            "summary": {"match": 1, "count_diff": 1, "left_only": 1, "right_only": 0},
            "gate_decision": {"policy": "auto", "passed": True},
            "timing": {"left_count_s": 1.5, "right_count_s": 0.5},
        },
    )
    store.put_json("manifest/p.json", {"pair": "p", "vintage": "month", "expected_units": 2})
    store.put_json("manifest/p.keys.json", {"pair": "p", "windows": {}})
    store.put_json("manifest/_CLOSED", {"pairs": ["p"], "expected_units": 2})
    same = compare_statistics(pair="p", window="2026-01-01", column="x", left=LEFT, right=LEFT)
    same["unit_id"], same["right_col"] = "u1", "x"
    changed = compare_statistics(
        pair="p",
        window="2026-01-01",
        column="y",
        left=LEFT,
        right={**LEFT, "mean": 9.0, "n_unique": 2},
    )
    changed["unit_id"], changed["right_col"] = "u2", "y_r"
    store.put_result("u1", same)
    store.put_result("u2", changed)
    return store


def test_build_report_ignores_left_documents_and_key_files(tmp_path):
    files = build_report(_store(), tmp_path / "out", title="T")
    assert set(files) == {
        "p_compare_row.csv",
        "p_compare_col.csv",
        "index.html",
        "report.xlsx",
        "summary.json",
    }
    summary = json.loads(files["summary.json"].read_text())
    assert summary["status"] == "DONE" and summary["pairs"] == 1
    assert summary["findings"] == {"row_count_diff": 1, "partition_missing": 1, "col_stat_diff": 1}
    page = files["index.html"].read_text()
    assert "days mismatch: 1" in page and "y → y_r" in page and "mean (2→9)" in page
    from openpyxl import load_workbook

    workbook = load_workbook(files["report.xlsx"])
    assert workbook.sheetnames == ["p_row", "p_col"]
    assert (
        workbook["p_col"]["A1"].value == "Vintage: 2026-01-01"
        and workbook["p_col"]["B1"].value == "FAIL"
    )


def test_report_files_can_be_fetched_from_a_run_store(tmp_path):
    files = build_report(_store(), tmp_path / "out")
    runs_store = VolumeStore(tmp_path / "runs")
    run_store = runs_store.child("run")
    for name, path in files.items():
        run_store.put_bytes(f"report/{name}", path.read_bytes(), content_type="")
    run_store.put_json("_STATUS.json", {"state": "complete"})
    fetched = fetch_report(runs_store, tmp_path / "fetched")
    assert fetched["run_id"] == "run"
    assert sorted(fetched["files"]) == sorted(files)
    assert (tmp_path / "fetched" / "report.xlsx").read_bytes() == files["report.xlsx"].read_bytes()


def test_compare_col_csv_orders_rows_and_unions_fields():
    rows = [
        {"vintage": "2026-02-01", "column": "a", "verdict": "match"},
        {"vintage": "2026-01-01", "column": "b", "extra": 1},
    ]
    assert compare_col_csv_bytes(rows).decode().splitlines() == [
        "column,extra,verdict,vintage",
        "b,1,,2026-01-01",
        "a,,match,2026-02-01",
    ]


def test_run_is_partial_with_dead_units_or_without_a_closed_manifest(tmp_path):
    store = _store()
    store.put_dead("u3", {"unit_id": "u3", "error": "boom"})
    summary = json.loads(build_report(store, tmp_path / "a")["summary.json"].read_text())
    assert summary["status"] == "PARTIAL" and summary["errors"] == [
        {"unit_id": "u3", "error": "boom"}
    ]
    store = _store()
    store.objects.pop("manifest/_CLOSED")
    summary = json.loads(build_report(store, tmp_path / "b")["summary.json"].read_text())
    assert summary["status"] == "PARTIAL" and summary["units_expected"] == 2


def test_a_run_without_col_results_still_reports_row_counts(tmp_path):
    store = _store()
    for key in [key for key in store.objects if key.startswith("results/")]:
        store.objects.pop(key)
    files = build_report(store, tmp_path)
    assert "p_compare_col.csv" not in files and "Stage 2" not in files["index.html"].read_text()


def test_build_report_needs_at_least_one_row_check_document(tmp_path):
    with pytest.raises(ValueError, match="no row_check documents"):
        build_report(MemoryStore(), tmp_path)
