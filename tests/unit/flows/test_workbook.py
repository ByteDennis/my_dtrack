import json
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from pair_recon.core.metadata.workbook import (
    FIELD_GUIDE_COLUMNS,
    PAIR_SQL_WORKBOOK_COLUMNS,
    PAIR_WORKBOOK_COLUMNS,
    changed_pair_configurations,
    create_pair_workbook,
    pair_configuration_signature,
    read_pair_workbook,
    upsert_pair_configuration,
)


class Sql:
    def __init__(self):
        self.calls = []

    def execute(self, statement, parameters=None):
        self.calls.append((statement, parameters or {}))
        return []


def set_row(sheet, values):
    for column, value in enumerate(values, 1):
        sheet.cell(2, column, value)


def valid_pair_row(*, keys="ORDER_ID", defaults='{"sample":0.1}'):
    return [
        "orders",
        True,
        keys,
        "month",
        "oracle",
        "PCDS",
        "SALES.ORDERS",
        "ORDER_DT",
        None,
        "databricks_query",
        None,
        None,
        "order_dt",
        None,
        defaults,
    ]


def test_create_and_read_pair_workbook(tmp_path):
    path = create_pair_workbook(tmp_path / "pairs.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["pairs", "pair_sql", "field_guide"]
    assert list(workbook["pairs"].tables) == ["PairReconPairs"]
    assert list(workbook["pair_sql"].tables) == ["PairReconSql"]
    assert list(workbook["field_guide"].tables) == []
    guide = {row[1]: row for row in workbook["field_guide"].iter_rows(min_row=2, values_only=True)}
    assert guide["run_defaults_json"][4:] == (
        "Run defaults; prefer sample fractions.",
        '{"sample":0.1}',
    )
    set_row(workbook["pairs"], valid_pair_row())
    set_row(
        workbook["pair_sql"],
        [
            "orders",
            "SELECT * FROM SALES.ORDERS l WHERE {{query_windows:driver}}",
            '{"driver":{"expression":"l.ORDER_DT"}}',
            "SELECT * FROM main.sales.orders r WHERE {{query_windows:driver}}",
            '{"driver":{"expression":"r.order_dt"}}',
        ],
    )
    workbook.save(path)

    rows = read_pair_workbook(path)

    assert rows == [
        {
            "pair_name": "orders",
            "enabled": True,
            "key_columns": ["ORDER_ID"],
            "vintage": "month",
            "left_source": "oracle",
            "left_conn": "PCDS",
            "left_table": "SALES.ORDERS",
            "left_date_col": "ORDER_DT",
            "left_options_json": None,
            "right_source": "databricks_query",
            "right_conn": None,
            "right_table": None,
            "right_date_col": "order_dt",
            "right_options_json": None,
            "run_defaults_json": '{"sample":0.1}',
            "pair_sql": {
                "left_sql": "SELECT * FROM SALES.ORDERS l WHERE {{query_windows:driver}}",
                "left_query_windows_json": '{"driver":{"expression":"l.ORDER_DT"}}',
                "right_sql": "SELECT * FROM main.sales.orders r WHERE {{query_windows:driver}}",
                "right_query_windows_json": '{"driver":{"expression":"r.order_dt"}}',
            },
        }
    ]


def test_pair_workbook_percentage_sample_requires_keys(tmp_path):
    path = create_pair_workbook(tmp_path / "pairs.xlsx")
    workbook = load_workbook(path)
    set_row(workbook["pairs"], valid_pair_row(keys=""))
    set_row(workbook["pair_sql"], ["orders", None, None, "SELECT 1", None])
    workbook.save(path)

    with pytest.raises(ValueError, match="sampled runs require key_columns"):
        read_pair_workbook(path)


FULL_RUN_DEFAULTS = {
    "run_mode": "full",
    "partition_mode": "date",
    "sample": 0.1,
    "strict": False,
    "checks": ["row_count", "hash_sum"],
    "audit": True,
    "sync_remote": False,
    "left_workers": 4,
    "right_workers": 16,
    "vintage": "month",
    "from_date": "2026-01-01",
    "to_date": "2026-02-28",
    "excluded_partitions": ["2026-01-02", "2026-02-01->2026-02-03"],
    "vintage_sample_size": 3,
}


def _read_with_defaults(tmp_path, defaults):
    path = create_pair_workbook(tmp_path / "pairs.xlsx")
    workbook = load_workbook(path)
    set_row(workbook["pairs"], valid_pair_row(defaults=json.dumps(defaults)))
    set_row(workbook["pair_sql"], ["orders", None, None, "SELECT 1", None])
    workbook.save(path)
    return read_pair_workbook(path)


def test_run_defaults_accept_every_prompt_default(tmp_path):
    rows = _read_with_defaults(tmp_path, FULL_RUN_DEFAULTS)
    assert json.loads(rows[0]["run_defaults_json"]) == FULL_RUN_DEFAULTS


@pytest.mark.parametrize(
    ("override", "message"),
    [
        ({"from_date": "2026/01/01"}, "from_date must be YYYY-MM-DD"),
        ({"to_date": "2026-1-5"}, "to_date must be YYYY-MM-DD"),
        ({"from_date": "2026-03-01"}, "from_date must not be after to_date"),
        ({"vintage": "weekly"}, "vintage must be one of day, week, month, quarter, year"),
        ({"audit": "yes"}, "audit must be true or false"),
        ({"sync_remote": 1}, "sync_remote must be true or false"),
        ({"left_workers": 0}, "left_workers must be a positive integer"),
        ({"right_workers": True}, "right_workers must be a positive integer"),
        ({"excluded_partitions": "2026-01-02"}, "excluded_partitions must be a string array"),
        ({"cadence": "x"}, "run_defaults_json has unsupported fields: cadence"),
    ],
)
def test_run_defaults_reject_bad_prompt_defaults(tmp_path, override, message):
    with pytest.raises((ValueError, TypeError), match=message):
        _read_with_defaults(tmp_path, {**FULL_RUN_DEFAULTS, **override})


def test_run_defaults_all_mode_rejects_date_only_fields(tmp_path):
    defaults = {"partition_mode": "all", "from_date": "2026-01-01", "vintage": "day"}
    with pytest.raises(ValueError, match="all mode cannot set from_date, vintage"):
        _read_with_defaults(tmp_path, defaults)


def test_pair_sql_must_reference_a_pair(tmp_path):
    path = create_pair_workbook(tmp_path / "pairs.xlsx")
    workbook = load_workbook(path)
    set_row(workbook["pairs"], valid_pair_row(defaults='{"partition_mode":"all"}'))
    set_row(workbook["pair_sql"], ["missing", None, None, "SELECT 1", None])
    workbook.save(path)

    with pytest.raises(ValueError, match="unknown pair_name 'missing'"):
        read_pair_workbook(path)


def test_upsert_pair_configuration_preserves_runtime_columns():
    sql = Sql()
    rows = [
        {
            "pair_name": "orders",
            "enabled": True,
            "key_columns": ["ORDER_ID"],
            "vintage": "month",
            "left_source": "oracle",
            "left_conn": "PCDS",
            "left_table": "SALES.ORDERS",
            "left_date_col": "ORDER_DT",
            "left_options_json": None,
            "right_source": "databricks_query",
            "right_conn": None,
            "right_table": None,
            "right_date_col": "order_dt",
            "right_options_json": None,
            "run_defaults_json": '{"sample":0.1}',
            "pair_sql": {
                "left_sql": None,
                "left_query_windows_json": None,
                "right_sql": "SELECT 1",
                "right_query_windows_json": None,
            },
        }
    ]

    result = upsert_pair_configuration(sql, "main.recon.pair_recon_pair_meta", rows)

    assert result["rows"] == result["sql_rows"] == 1
    assert result["pair_sql_table"] == "main.recon.pair_recon_pair_sql"
    assert len(sql.calls) == 2
    meta_sql, meta_values = sql.calls[0]
    pair_sql, pair_values = sql.calls[1]
    assert "latest_run_id, latest_status" in meta_sql
    assert "t.latest_run_id" not in meta_sql and "t.latest_status" not in meta_sql
    assert "t.left_source = :left_source" in meta_sql
    assert "MERGE INTO main.recon.pair_recon_pair_sql" in pair_sql
    assert meta_values["enabled"] == "true"
    assert json.loads(meta_values["key_columns"]) == ["ORDER_ID"]
    assert json.loads(meta_values["config_json"])["left"]["conn_macro"] == "PCDS"
    assert pair_values["right_sql"] == "SELECT 1"


def test_pair_configuration_comparison_is_semantic_and_detects_changes():
    row = {
        "pair_name": "orders",
        "enabled": True,
        "key_columns": ["ORDER_ID"],
        "vintage": "month",
        "left_source": "oracle",
        "left_conn": "PCDS",
        "left_table": "SALES.ORDERS",
        "left_date_col": "ORDER_DT",
        "left_options_json": '{"where":"ACTIVE=1","date_type":"date"}',
        "right_source": "databricks_query",
        "right_conn": None,
        "right_table": None,
        "right_date_col": "order_dt",
        "right_options_json": None,
        "run_defaults_json": '{"sample":0.1}',
        "pair_sql": {
            "left_sql": None,
            "left_query_windows_json": None,
            "right_sql": "SELECT 1",
            "right_query_windows_json": None,
        },
    }
    remote = deepcopy(row)
    remote.update(
        enabled="true",
        key_columns='["ORDER_ID"]',
        left_options_json='{ "date_type": "date", "where": "ACTIVE=1" }',
        run_defaults_json='{ "sample": 0.1 }',
    )
    remote["pair_sql"]["updated_at"] = "2026-10-21T10:00:00"

    assert pair_configuration_signature(row) == pair_configuration_signature(remote)
    assert changed_pair_configurations([row], [remote]) == []

    remote["config_json"] = None
    assert changed_pair_configurations([row], [remote]) == [row]
    remote.pop("config_json")
    remote["left_table"] = "SALES.OLD_ORDERS"
    assert changed_pair_configurations([row], [remote]) == [row]
    assert changed_pair_configurations([row], []) == [row]


def test_pair_workbook_columns_are_stable():
    assert PAIR_WORKBOOK_COLUMNS == (
        "pair_name",
        "enabled",
        "key_columns",
        "vintage",
        "left_source",
        "left_conn",
        "left_table",
        "left_date_col",
        "left_options_json",
        "right_source",
        "right_conn",
        "right_table",
        "right_date_col",
        "right_options_json",
        "run_defaults_json",
    )
    assert PAIR_SQL_WORKBOOK_COLUMNS == (
        "pair_name",
        "left_sql",
        "left_query_windows_json",
        "right_sql",
        "right_query_windows_json",
    )
    assert FIELD_GUIDE_COLUMNS == (
        "worksheet",
        "field_name",
        "required_when",
        "type_default",
        "description",
        "example",
    )
