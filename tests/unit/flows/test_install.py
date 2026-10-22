import importlib.util
import json
from pathlib import Path
from types import SimpleNamespace as NS

import pytest

ROOT = Path(__file__).resolve().parents[3]
SPEC = importlib.util.spec_from_file_location("pair_recon_setup_install", ROOT / "setup/install.py")
assert SPEC and SPEC.loader
install = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(install)


def item(path, kind):
    return NS(path=path, object_type=NS(value=kind))


class WorkspaceFiles:
    def __init__(self, objects):
        self.objects = {value.path: value for value in objects}
        self.deleted = []

    def mkdirs(self, _path):
        return None

    def upload(self, path, _body, **options):
        kind = "NOTEBOOK" if options["format"].value == "SOURCE" else "FILE"
        self.objects[path] = item(path, kind)

    def list(self, path):
        prefix = path.rstrip("/") + "/"
        return iter(
            value
            for name, value in sorted(self.objects.items())
            if name.startswith(prefix) and "/" not in name[len(prefix) :]
        )

    def delete(self, path, *, recursive=None):
        self.deleted.append((path, recursive))
        prefix = path.rstrip("/") + "/"
        for name in list(self.objects):
            if name == path or (recursive and name.startswith(prefix)):
                del self.objects[name]


def test_upload_removes_remote_files_missing_locally(tmp_path):
    source = tmp_path / "pair_recon_right"
    source.mkdir()
    (source / "notebook.py").write_text("print('notebook')\n", encoding="utf-8")
    (source / "keep.py").write_text("keep = True\n", encoding="utf-8")
    (source / "new.py").write_text("new = True\n", encoding="utf-8")
    root = "/Users/me@example.com/app/pair_recon_right"
    files = WorkspaceFiles(
        [
            item(f"{root}/notebook", "NOTEBOOK"),
            item(f"{root}/keep.py", "FILE"),
            item(f"{root}/old.py", "FILE"),
            item(f"{root}/old", "DIRECTORY"),
            item(f"{root}/old/child.py", "FILE"),
        ]
    )

    uploaded, deleted = install.upload(NS(workspace=files), source, root)

    assert uploaded == [f"{root}/keep.py", f"{root}/new.py", f"{root}_job"]
    assert set(deleted) == {
        f"{root}/notebook",
        f"{root}/old.py",
        f"{root}/old",
        f"{root}/old/child.py",
    }
    assert set(files.objects) == {
        f"{root}/keep.py",
        f"{root}/new.py",
        f"{root}_job",
    }


def test_upload_failure_does_not_delete_remote_files(tmp_path):
    source = tmp_path / "pair_recon_right"
    source.mkdir()
    (source / "new.py").write_text("new = True\n", encoding="utf-8")
    root = "/Users/me@example.com/app/pair_recon_right"
    files = WorkspaceFiles([item(f"{root}/old.py", "FILE")])

    def fail_upload(*_args, **_kwargs):
        raise OSError("upload failed")

    files.upload = fail_upload
    with pytest.raises(OSError, match="upload failed"):
        install.upload(NS(workspace=files), source, root)

    assert files.deleted == []
    assert f"{root}/old.py" in files.objects


def test_refresh_configuration_uses_workbook_and_statement_sql(monkeypatch, tmp_path):
    workbook = tmp_path / "pairs.xlsx"
    workbook.touch()
    rows = [{"pair_name": "orders"}]
    sql = object()
    monkeypatch.setattr(
        install, "read_pair_workbook", lambda path: rows if path == workbook else []
    )
    monkeypatch.setattr(install, "StatementSql", lambda workspace, warehouse: sql)
    monkeypatch.setattr(
        install,
        "upsert_pair_configuration",
        lambda actual_sql, table, actual_rows: {
            "sql": actual_sql,
            "table": table,
            "rows": actual_rows,
        },
    )

    result = install.refresh_configuration(
        object(), "warehouse", "main.recon.pair_recon_pair_meta", workbook
    )

    assert result == {
        "sql": sql,
        "table": "main.recon.pair_recon_pair_meta",
        "rows": rows,
    }


def test_setup_assets_are_three_sheet_workbook_and_mapping_notebook():
    from openpyxl import load_workbook

    workbook = load_workbook(ROOT / "setup" / "pairs.xlsx")
    assert workbook.sheetnames == ["pairs", "pair_sql", "field_guide"]
    assert list(workbook["pairs"].tables) == ["PairReconPairs"]
    assert list(workbook["pair_sql"].tables) == ["PairReconSql"]
    assert list(workbook["field_guide"].tables) == []
    notebook = json.loads((ROOT / "setup" / "mapping.ipynb").read_text(encoding="utf-8"))
    source = "\n".join(line for cell in notebook["cells"] for line in cell.get("source", []))
    assert "discover_oracle_columns" in source
    assert "discover_databricks_columns" in source
    assert "upsert_column_mappings" in source
    assert "load_environment()" in source
    assert "DBX_WAREHOUSE_ID" in source
    assert "pair_config(row)" in source
    assert "APPLY = False" in source

    right_notebook = (ROOT / "src/pair_recon_right/notebook.py").read_text(encoding="utf-8")
    assert right_notebook.count("**publish_values") == 2
    assert "publish(spark, store, **values" not in right_notebook
    assert "from pair_recon_right.catalog import publish" in right_notebook
    assert "importlib" not in right_notebook
    assert "sys.path" not in right_notebook
    assert "force=True" in right_notebook
    assert "stream=sys.stdout" in right_notebook
    assert right_notebook.rstrip().endswith(
        "dbutils.notebook.exit(json.dumps(outcome, default=str))"
    )
    assert right_notebook.index("except BaseException") < right_notebook.index(
        "dbutils.notebook.exit"
    )

    schema = (ROOT / "setup" / "schema.sql").read_text(encoding="utf-8")
    pair_sql_schema = schema.split("pair_recon_pair_sql", 1)[1].split(");", 1)[0]
    assert "left_source STRING NOT NULL" in schema
    assert "run_defaults_json STRING" in schema
    assert "left_query_windows_json STRING" in pair_sql_schema
    assert "right_query_windows_json STRING" in pair_sql_schema
    assert "vintage_sample_size" not in pair_sql_schema
    assert "row_differences BIGINT" in schema
    assert "column_differences BIGINT" in schema
    assert "runtime_failures BIGINT" in schema
    migration = (ROOT / "setup" / "migrate_pair_metadata.sql").read_text(encoding="utf-8")
    assert "left_conn STRING" in migration
    assert "right_query_windows_json STRING" in migration
    status_migration = (ROOT / "setup" / "migrate_run_status.sql").read_text(encoding="utf-8")
    assert "row_differences BIGINT" in status_migration
    assert "column_differences BIGINT" in status_migration
    assert "runtime_failures BIGINT" in status_migration


def test_main_uses_dbx_environment_defaults(monkeypatch, tmp_path, capsys):
    monkeypatch.chdir(tmp_path)
    for name in (
        "DBX_PROFILE",
        "DBX_WORKSPACE_PDIR",
        "DBX_CLUSTER_ID",
        "DBX_WAREHOUSE_ID",
        "DBX_PAIR_TABLE",
        "DBX_PAIR_WORKBOOK",
    ):
        monkeypatch.delenv(name, raising=False)
    (tmp_path / ".pair_recon.env").write_text(
        "DBX_PROFILE=team\nDBX_WORKSPACE_PDIR=deploy\nDBX_CLUSTER_ID=c9\n"
        "DBX_WAREHOUSE_ID=w9\nDBX_PAIR_TABLE=main.recon.pair_recon_pair_meta\n"
        "DBX_PAIR_WORKBOOK=config/pairs.xlsx\n",
        encoding="utf-8",
    )
    client = object()
    calls = {}
    monkeypatch.setattr(
        install,
        "WorkspaceClient",
        lambda *, profile: calls.setdefault("profile", profile) and client,
    )
    monkeypatch.setattr(
        install,
        "right_target",
        lambda actual, path: (
            calls.update({"resolve": (actual, path)}) or f"/Users/me/{path}/pair_recon"
        ),
    )
    monkeypatch.setattr(install, "read_pair_workbook", lambda _path: [])
    monkeypatch.setattr(install, "upload", lambda *_args: (["uploaded.py"], []))
    monkeypatch.setattr(
        install, "upsert", lambda actual, spec: calls.update({"job": (actual, spec)}) or 7
    )
    monkeypatch.setattr(
        install,
        "refresh_configuration",
        lambda actual, warehouse, table, workbook: (
            calls.update({"refresh": (actual, warehouse, table, workbook)}) or {"pairs": 0}
        ),
    )

    assert install.main([]) == 0

    assert calls["profile"] == "team"
    assert calls["resolve"] == (client, "deploy")
    assert calls["job"][1]["tasks"][0]["notebook_task"]["notebook_path"] == (
        "/Users/me/deploy/pair_recon_job"
    )
    assert calls["job"][1]["tasks"][0]["existing_cluster_id"] == "c9"
    assert calls["refresh"][1:3] == ("w9", "main.recon.pair_recon_pair_meta")
    assert calls["refresh"][3] == Path("config/pairs.xlsx")
    assert json.loads(capsys.readouterr().out)["job_id"] == 7
