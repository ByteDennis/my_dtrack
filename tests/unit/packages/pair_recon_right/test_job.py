import pytest
from support import FakeTarget, MemoryStore

from pair_recon_right.engine import job
from pair_recon_right.engine.backends.store import VolumeStore
from pair_recon_right.engine.job import run, summary


def document(checks):
    side = {
        "left": {"source": "oracle"},
        "right": {"source": "databricks", "table": "R"},
    }
    return {
        "run_id": "r1",
        "pair": "p",
        **side,
        "bounds": {"from": None, "to": None},
        "partition_mode": "all",
        "counts": {"all": 2},
        "pair_config": {
            "version": 2,
            "partition_mode": "all",
            **side,
            "col_map": {"A": "a"},
            "checks": checks,
        },
        "health_check": {"ok": True},
        "timing": {},
    }


def test_rowonly(tmp_path):
    store = VolumeStore(tmp_path / "runs" / "r1")
    store.put_json("row_check/left.json", document(["row_count"]))
    store.put_json("_status/left_row_check.closed.json", {"pairs": ["p"]})
    store.put_bytes("report/row_check.csv", b"old")
    store.put_bytes("report/col_check.csv", b"old")
    result = run(
        tmp_path / "runs",
        lambda _config: FakeTarget(counts={"all": 2}),
        run_id="r1",
        poll_interval=0,
    )
    assert result == {"row_errors": 0, "column_errors": 0, "failures": 0}
    assert set(store.list_keys("report/")) == {
        "report/index.html",
        "report/report.xlsx",
        "report/summary.json",
    }


def test_run_waits_for_left_row_data_without_a_deadline(tmp_path, monkeypatch):
    store = VolumeStore(tmp_path / "runs" / "r1")
    polls = []

    def sleep(_seconds):
        polls.append(1)
        if len(polls) == 3:
            store.put_json("row_check/left.json", document(["row_count"]))
            store.put_json("_status/left_row_check.closed.json", {"pairs": ["p"]})

    monkeypatch.setattr(job.time, "sleep", sleep)
    result = run(
        tmp_path / "runs",
        lambda _config: FakeTarget(counts={"all": 2}),
        run_id="r1",
        poll_interval=0,
    )
    assert result["row_errors"] == 0 and len(polls) == 3


def test_run_stops_waiting_when_the_left_heartbeat_is_stale(tmp_path):
    store = VolumeStore(tmp_path / "runs" / "r1")
    store.heartbeat(now=0)
    with pytest.raises(RuntimeError, match="heartbeat is stale"):
        run(tmp_path / "runs", None, run_id="r1", poll_interval=0, heartbeat_timeout=10)


def test_run_refuses_a_plan_without_units(tmp_path):
    store = VolumeStore(tmp_path / "runs" / "r1")
    store.put_json("row_check/left.json", document(["hash_sum"]))
    store.put_json("_status/left_row_check.closed.json", {"pairs": ["p"]})
    store.put_json(
        "plan.json",
        {"pair": "p", "version": 2, "windows": [], "columns": ["A"], "expected_units": 0},
    )
    with pytest.raises(RuntimeError, match="plan has no units to compare"):
        run(
            None,
            lambda _config: FakeTarget(counts={"all": 2}),
            run_id="r1",
            store=store,
            poll_interval=0,
        )


def test_row_check_failure_publishes_partial_workbook_without_csv(tmp_path):
    store = VolumeStore(tmp_path / "runs" / "r1")
    store.put_json("row_check/left.json", document(["hash_sum"]))
    store.put_json("_status/left_row_check.closed.json", {"pairs": ["p"]})

    with pytest.raises(RuntimeError, match="partition_mode all row counts differ"):
        run(
            None,
            lambda _config: FakeTarget(counts={"all": 3}),
            run_id="r1",
            poll_interval=0,
            store=store,
        )

    assert set(store.list_keys("report/")) == {
        "report/index.html",
        "report/report.xlsx",
        "report/summary.json",
    }

    from openpyxl import load_workbook

    assert load_workbook(tmp_path / "runs" / "r1" / "report" / "report.xlsx").sheetnames == [
        "Row Check"
    ]


def test_summary_excludes_stale_outcomes_from_an_empty_current_plan():
    store = MemoryStore()
    store.put_json("plan.json", {"pair": "p", "windows": [], "columns": []})
    store.put_json(
        "col_check/results/stale.json",
        {"unit_id": "stale", "pair": "p", "verdict": "diff"},
    )
    store.put_json(
        "col_check/failures/stale.json",
        {"unit_id": "stale-failure", "pair": "p", "error": "old"},
    )

    assert summary(store) == {"row_errors": 0, "column_errors": 0, "failures": 0}
