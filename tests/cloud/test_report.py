import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from conftest import MemoryStore
from dtrack_cloud.comparison import compare_statistics
from dtrack_cloud.reporting import build_report, publish_report

from dtrack_distributed.workflows.fetch import fetch_report

LEFT = {
    "col_type": "numeric",
    "n_total": 3,
    "n_missing": 0,
    "n_unique": 3,
    "hash_sum": "10",
    "mean": 2.0,
    "std": 1.0,
    "min_val": "1",
    "max_val": "3",
    "length_avg": None,
    "length_max": None,
    "top_10": "",
}


def _store():
    store = MemoryStore()
    store.put_json("stage1/p.left.json", {"pair": "p", "counts": {}})
    store.put_json(
        "stage1/p.json",
        {
            "run_id": "run",
            "pair": "p",
            "left": {"source": "oracle", "table": "A", "date_col": "DT"},
            "right": {
                "source": "databricks",
                "catalog": "c",
                "schema": "s",
                "table": "a",
                "date_col": "dt",
            },
            "bounds": {"from": "2026-01-01", "to": "2026-01-03"},
            "partitions": {
                "2026-01-01": {"left": 1, "right": 1, "verdict": "match"},
                "2026-01-02": {"left": 1, "right": 2, "verdict": "count_diff", "delta": 1},
                "2026-01-03": {"left": 1, "right": None, "verdict": "left_only"},
            },
            "summary": {"match": 1, "count_diff": 1, "left_only": 1, "right_only": 0},
            "gate_decision": {"policy": "auto", "passed": True},
            "timing": {"left_count_s": 1.5, "right_count_s": 0.5},
        },
    )
    store.put_json("manifest/p.json", {"pair": "p", "vintage": "month", "expected_units": 2})
    store.put_json("manifest/p.keys.json", {"pair": "p", "windows": {}})
    store.put_json("manifest/_CLOSED", {"pairs": ["p"], "expected_units": 2})
    same = compare_statistics(pair="p", window="2026-01-01", column="x", left=LEFT, right=LEFT)
    same["unit_id"], same["right_col"] = "u1", "x"
    changed = compare_statistics(
        pair="p",
        window="2026-01-01",
        column="y",
        left=LEFT,
        right={**LEFT, "mean": 9.0, "n_unique": 2},
    )
    changed["unit_id"], changed["right_col"] = "u2", "y_r"
    store.put_result("u1", same)
    store.put_result("u2", changed)
    return store


def test_build_report_ignores_left_documents_and_key_files(tmp_path):
    files = build_report(_store(), tmp_path / "out", title="T")
    assert set(files) == {
        "p_compare_row.csv",
        "p_compare_col.csv",
        "index.html",
        "report.xlsx",
        "summary.json",
    }
    summary = json.loads(files["summary.json"].read_text())
    assert summary["status"] == "DONE" and summary["pairs"] == 1
    assert summary["findings"] == {"row_count_diff": 1, "partition_missing": 1, "col_stat_diff": 1}
    page = files["index.html"].read_text()
    assert "days mismatch: 1" in page and "y → y_r" in page and "mean (2→9)" in page
    from openpyxl import load_workbook

    workbook = load_workbook(files["report.xlsx"])
    assert workbook.sheetnames == ["p_row", "p_col"]
    assert (
        workbook["p_col"]["A1"].value == "Vintage: 2026-01-01"
        and workbook["p_col"]["B1"].value == "FAIL"
    )


def test_publish_then_fetch_round_trips_the_report_set(tmp_path):
    files = build_report(_store(), tmp_path / "out")
    summary = json.loads(files["summary.json"].read_text())
    report_store = MemoryStore()
    keys = publish_report(report_store, "run", files, summary)
    assert keys[-1] == "latest/pointer.json" and report_store.events[-1] == (
        "put",
        "latest/pointer.json",
    )
    fetched = fetch_report(report_store, tmp_path / "fetched")
    assert fetched["run_id"] == "run"
    assert sorted(fetched["files"]) == sorted(files)
    assert (tmp_path / "fetched" / "report.xlsx").read_bytes() == files["report.xlsx"].read_bytes()
