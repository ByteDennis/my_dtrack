

# >>> build the row-comparison workbook (one sheet per pair) — see web /api/compare/row/export/excel-all <<< #
def build_row_comparison_workbook(db_path, config_path, from_date=None, to_date=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from .db import list_table_pairs, get_metadata
    from .compare import compare_row_counts
    from .config import load_unified_config

    config = load_unified_config(config_path)
    pairs = list_table_pairs(db_path)

    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    red_font = Font(color="CC0000")
    bold = Font(bold=True)

    wb = Workbook()
    wb.remove(wb.active)

    for pair in pairs:
        pname = pair["pair_name"]
        pair_cfg = config.get("pairs", {}).get(pname, {})
        if pair_cfg.get("skip"):
            continue

        comparison = compare_row_counts(
            db_path, pair["table_left"], pair["table_right"],
            from_date=from_date or None, to_date=to_date or None,
        )

        meta_left = get_metadata(db_path, pair["table_left"]) or {}
        meta_right = get_metadata(db_path, pair["table_right"]) or {}
        source_left = (pair.get("source_left") or "LEFT").upper()
        source_right = (pair.get("source_right") or "RIGHT").upper()
        table_left = meta_left.get("source_table") or pair["table_left"]
        table_right = meta_right.get("source_table") or pair["table_right"]
        date_var = meta_left.get("date_var") or meta_right.get("date_var") or ""
        date_fmt = meta_left.get("date_format") or meta_right.get("date_format") or ""

        ws = wb.create_sheet(title=pname[:31])

        s = comparison["summary"]
        overlap_start = overlap_end = None
        lmin, lmax = s["date_range_left"]
        rmin, rmax = s["date_range_right"]
        if lmin and rmin and lmax and rmax:
            overlap_start = max(lmin, rmin)
            overlap_end = min(lmax, rmax)

        ws.cell(row=1, column=1, value=f"{source_left}:").font = bold
        ws.cell(row=1, column=2, value=table_left)
        ws.cell(row=2, column=1, value=f"{source_right}:").font = bold
        ws.cell(row=2, column=2, value=table_right)
        date_label = f"{date_var} ({date_fmt})" if date_fmt else date_var
        ws.cell(row=3, column=1, value="Date Variable:").font = bold
        ws.cell(row=3, column=2, value=date_label)
        ws.cell(row=4, column=1, value="Overlap:").font = bold
        ws.cell(row=4, column=2, value=f"{overlap_start or '?'} -> {overlap_end or '?'}")
        ws.cell(row=5, column=1, value="Matching:").font = bold
        ws.cell(row=5, column=2, value=f"{s['count_left']} dates" if s.get('count_left') else "0")
        ws.cell(row=6, column=1, value="Mismatch:").font = bold
        ws.cell(row=6, column=2, value=len(comparison["mismatched"]))
        ws.cell(row=7, column=1, value=f"{source_left} Only:").font = bold
        ws.cell(row=7, column=2, value=len(comparison["only_left"]))
        ws.cell(row=8, column=1, value=f"{source_right} Only:").font = bold
        ws.cell(row=8, column=2, value=len(comparison["only_right"]))

        row_num = 10
        ws.cell(row=row_num, column=1, value="Date").font = bold
        ws.cell(row=row_num, column=2, value=source_left).font = bold
        ws.cell(row=row_num, column=3, value=source_right).font = bold
        ws.cell(row=row_num, column=4, value="Status").font = bold
        row_num += 1

        dates = []
        for dt, l, r in comparison["mismatched"]:
            dates.append((0, dt, l, r, "mismatch"))
        for dt, count in comparison["only_left"]:
            dates.append((0, dt, count, None, f"{source_left.lower()}_only"))
        for dt, count in comparison["only_right"]:
            dates.append((0, dt, None, count, f"{source_right.lower()}_only"))
        for dt, count in comparison["matching"]:
            dates.append((1, dt, count, count, "match"))
        dates.sort(key=lambda r: (r[0], r[1]))

        for _sort_key, dt, left, right, status in dates:
            ws.cell(row=row_num, column=1, value=dt)
            ws.cell(row=row_num, column=2, value=left)
            ws.cell(row=row_num, column=3, value=right)
            ws.cell(row=row_num, column=4, value=status)
            if status != "match":
                for col in range(1, 5):
                    cell = ws.cell(row=row_num, column=col)
                    cell.fill = red_fill
                    cell.font = red_font
            row_num += 1

    return wb

# >>> build the col-comparison workbook (one sheet per pair, transposed) — see web /api/compare/col/export/excel <<< #
def build_col_comparison_workbook(db_path, config_path, from_date=None,
                                  to_date=None, compare_fields=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from .db import (
        list_table_pairs, get_metadata, get_pair_col_map_from_db, get_column_meta,
    )
    from .config import load_unified_config, get_col_type_overrides
    from .compare import compare_column_stats, _has_col_differences, match_columns_from_dicts

    cf = compare_fields
    config = load_unified_config(config_path)
    pairs = list_table_pairs(db_path)

    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    bold = Font(bold=True)

    wb = Workbook()
    wb.remove(wb.active)

    stat_rows_all = ["Type", "N_Total", "N_Missing", "N_Unique", "Mean", "Std",
                     "Min", "Max", "Length avg", "Length max", "Hash sum"]
    stat_keys_l_all = ["col_type", "n_total_left", "n_missing_left", "n_unique_left",
                       "mean_left", "std_left", "min_left", "max_left",
                       "length_avg_left", "length_max_left", "hash_sum_left"]
    stat_keys_r_all = ["col_type", "n_total_right", "n_missing_right", "n_unique_right",
                       "mean_right", "std_right", "min_right", "max_right",
                       "length_avg_right", "length_max_right", "hash_sum_right"]
    _ROW_TO_FIELD = {
        "Type": None, "N_Total": "n_total", "N_Missing": "n_missing",
        "N_Unique": "n_unique", "Mean": "mean", "Std": "std",
        "Min": "min", "Max": "max", "Length avg": "length_avg",
        "Length max": "length_max", "Hash sum": "hash_sum",
    }
    if cf is not None:
        keep_idx = [i for i, label in enumerate(stat_rows_all)
                    if _ROW_TO_FIELD[label] is None or _ROW_TO_FIELD[label] in cf]
        stat_rows_all = [stat_rows_all[i] for i in keep_idx]
        stat_keys_l_all = [stat_keys_l_all[i] for i in keep_idx]
        stat_keys_r_all = [stat_keys_r_all[i] for i in keep_idx]

    def _filter_active_stats(comps):
        def _has_value(key):
            for c in comps:
                v = c.get(key)
                if v is not None and v != "":
                    return True
            return False
        active = []
        for i in range(len(stat_rows_all)):
            if _has_value(stat_keys_l_all[i]) or _has_value(stat_keys_r_all[i]):
                active.append(i)
        return (
            [stat_rows_all[i] for i in active],
            [stat_keys_l_all[i] for i in active],
            [stat_keys_r_all[i] for i in active],
        )

    for pair in pairs:
        pname = pair["pair_name"]
        pair_cfg = config.get("pairs", {}).get(pname, {})
        if pair_cfg.get("skip"):
            continue

        col_type_overrides = get_col_type_overrides(config, pname)
        col_mappings = get_pair_col_map_from_db(db_path, pname)
        if not col_mappings:
            left_meta = get_column_meta(db_path, pair["table_left"])
            right_meta = get_column_meta(db_path, pair["table_right"])
            left_cols = {c["column_name"]: c.get("data_type", "") for c in left_meta}
            right_cols = {c["column_name"]: c.get("data_type", "") for c in right_meta}
            auto = match_columns_from_dicts(left_cols, right_cols)
            col_mappings = auto.get("matched", {})

        comparison = compare_column_stats(
            db_path, pair["table_left"], pair["table_right"],
            columns=list(col_mappings.keys()) if col_mappings else None,
            col_mappings=col_mappings,
            from_date=from_date or None, to_date=to_date or None,
            col_type_overrides=col_type_overrides,
        )

        meta_left = get_metadata(db_path, pair["table_left"]) or {}
        meta_right = get_metadata(db_path, pair["table_right"]) or {}
        source_left = (pair.get("source_left") or "LEFT").upper()
        source_right = (pair.get("source_right") or "RIGHT").upper()
        table_left = meta_left.get("source_table") or pair["table_left"]
        table_right = meta_right.get("source_table") or pair["table_right"]

        # original per-side column types (verbatim from columns_file / discovery,
        # e.g. athena 'varchar', hadoop 'string', oracle 'NUMBER') for the Type row
        left_type_map = {c["column_name"]: c.get("data_type", "")
                         for c in (get_column_meta(db_path, pair["table_left"]) or [])}
        right_type_map = {c["column_name"]: c.get("data_type", "")
                          for c in (get_column_meta(db_path, pair["table_right"]) or [])}

        vintages = {}
        for col_name, comps in comparison.items():
            for comp in comps:
                vl = comp.get("vintage_label") or comp.get("dt", "")
                vintages.setdefault(vl, []).append(comp)

        ws = wb.create_sheet(title=pname[:31])
        row_num = 1

        DIFF_BY_LABEL = {
            "N_Total":    lambda c: c.get("n_total_diff", 0) != 0,
            "N_Missing":  lambda c: c.get("n_missing_diff", 0) != 0,
            "N_Unique":   lambda c: c.get("n_unique_diff", 0) != 0,
            "Mean":       lambda c: c.get("mean_match") is False,
            "Std":        lambda c: c.get("std_match") is False,
            "Min":        lambda c: c.get("min_match") is False,
            "Max":        lambda c: c.get("max_match") is False,
            "Length avg": lambda c: c.get("length_avg_match") is False,
            "Length max": lambda c: c.get("length_max_match") is False,
            "Hash sum":   lambda c: c.get("hash_sum_match") is False,
        }

        for vl in sorted(vintages.keys()):
            comps = vintages[vl]
            has_diff = any(_has_col_differences(c, compare_fields=cf) for c in comps)

            comps.sort(key=lambda c: (0 if _has_col_differences(c, compare_fields=cf) else 1,
                                      c.get("left_col", "")))
            col_names_left = [c.get("left_col", "") for c in comps]
            col_names_right = [c.get("right_col", "") or c.get("left_col", "") for c in comps]

            stat_rows, stat_keys_l, stat_keys_r = _filter_active_stats(comps)

            diff_cells = set()
            problem_cols = set()
            for ci, comp in enumerate(comps):
                for si, label in enumerate(stat_rows):
                    check = DIFF_BY_LABEL.get(label)
                    if check is not None and check(comp):
                        diff_cells.add((ci, si))
            for ci, _si in diff_cells:
                problem_cols.add(ci)

            ws.cell(row=row_num, column=1, value=f"Vintage: {vl}").font = bold
            status_cell = ws.cell(row=row_num, column=2, value="FAIL" if has_diff else "PASS")
            status_cell.fill = red_fill if has_diff else green_fill
            status_cell.font = bold
            row_num += 1

            ws.cell(row=row_num, column=1, value=source_left).font = bold
            ws.cell(row=row_num, column=2, value=table_left)
            row_num += 1
            for ci, cn in enumerate(col_names_left):
                ws.cell(row=row_num, column=ci + 2, value=cn)
            row_num += 1
            for si, (stat_label, stat_key) in enumerate(zip(stat_rows, stat_keys_l)):
                ws.cell(row=row_num, column=1, value=stat_label).font = bold
                for ci, comp in enumerate(comps):
                    if stat_label == "Type":
                        lc = comp.get("left_col", "")
                        val = (left_type_map.get(lc) or comp.get("col_type_left")
                               or comp.get("col_type", ""))
                    else:
                        val = comp.get(stat_key, "")
                        if val is not None and val != "":
                            try:
                                val = float(val)
                            except (ValueError, TypeError):
                                pass
                    cell = ws.cell(row=row_num, column=ci + 2, value=val if val != "" else None)
                    if (ci, si) in diff_cells:
                        cell.fill = red_fill
                    elif ci in problem_cols and si > 0:
                        cell.fill = green_fill
                row_num += 1

            row_num += 1

            ws.cell(row=row_num, column=1, value=source_right).font = bold
            ws.cell(row=row_num, column=2, value=table_right)
            row_num += 1
            for ci, cn in enumerate(col_names_right):
                ws.cell(row=row_num, column=ci + 2, value=cn)
            row_num += 1
            for si, (stat_label, stat_key) in enumerate(zip(stat_rows, stat_keys_r)):
                ws.cell(row=row_num, column=1, value=stat_label).font = bold
                for ci, comp in enumerate(comps):
                    if stat_label == "Type":
                        rc = comp.get("right_col", "") or comp.get("left_col", "")
                        val = (right_type_map.get(rc) or comp.get("col_type_right")
                               or comp.get("col_type", ""))
                    else:
                        val = comp.get(stat_key, "")
                        if val is not None and val != "":
                            try:
                                val = float(val)
                            except (ValueError, TypeError):
                                pass
                    cell = ws.cell(row=row_num, column=ci + 2, value=val if val != "" else None)
                    if (ci, si) in diff_cells:
                        cell.fill = red_fill
                    elif ci in problem_cols and si > 0:
                        cell.fill = green_fill
                row_num += 1

            row_num += 2

    return wb
