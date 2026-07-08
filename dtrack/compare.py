
import json
import re
from fnmatch import fnmatch
from .db import (get_row_counts, get_col_stats, get_table_pair, list_table_pairs,
                 get_pair_col_map_from_db, get_column_meta)
from .config import load_unified_config
from .date_utils import parse_date
from .constants import STAT_ROUND_DECIMALS
from .platforms.base import (quote_ident, resolve_table, qualified_name,
                             _is_numeric_col_type, _is_date_col_type)

_DATE_SENTINELS = frozenset({'all', 'sample'})

# >>> normalize a col_type label (raw DB type or dtrack tag) to numeric/date/categorical <<< #
def _col_family(t):
    if _is_numeric_col_type(t):
        return 'numeric'
    if _is_date_col_type(t):
        return 'date'
    return 'categorical'

# >>> Apply include/exclude glob patterns against col_map keys <<< #
def resolve_col_filter(col_map, include_patterns=None, exclude_patterns=None):
    include_patterns = [p.strip() for p in (include_patterns or []) if p and p.strip()]
    exclude_patterns = [p.strip() for p in (exclude_patterns or []) if p and p.strip()]

    mapped_left_lower = {l.lower(): l for l in (col_map or {}).keys()}

    def _matches_any(name_lower, pats):
        return any(fnmatch(name_lower, p.lower()) for p in pats)

    if include_patterns:
        kept_lower = {
            k: orig for k, orig in mapped_left_lower.items()
            if _matches_any(k, include_patterns)
        }
    else:
        kept_lower = dict(mapped_left_lower)

    if exclude_patterns:
        kept_lower = {
            k: orig for k, orig in kept_lower.items()
            if not _matches_any(k, exclude_patterns)
        }

    kept = sorted(kept_lower.values(), key=str.lower)
    pairs = [(l, col_map[l]) for l in kept]

    unmapped_matches = []
    if include_patterns:
        unmapped_matches = []

    return {
        "pairs": pairs,
        "total_mapped": len(col_map or {}),
        "effective_count": len(pairs),
        "unmapped_matches": unmapped_matches,
    }

# >>> Parse a value to int, returning 0 for empty/None <<< #
def _safe_int(val):
    if val is None or val == '':
        return 0
    return int(float(str(val)))

# >>> Parse a value to float, returning None for empty/None <<< #
def _safe_float(val):
    if val is None or val == '':
        return None
    return float(str(val))

# >>> Build a string match key for mean/std/min/max <<< #
def _stat_match_key(val, prefix, ndigits=STAT_ROUND_DECIMALS):
    if val is None or val == '':
        return None
    s = str(val).strip()
    if prefix in ("mn", "mx") and "=" in s:
        s = s.rsplit("=", 1)[0].strip()
    looks_date = any(c.isalpha() or c in '-/: ' for c in s)
    if looks_date:
        try:
            from .date_utils import parse_date
            return f"{prefix}_{parse_date(s)}"
        except (ValueError, TypeError):
            pass
    try:
        return f"{prefix}_{round(float(s), ndigits)}"
    except (ValueError, TypeError):
        pass
    return f"{prefix}_{s}"

# >>> Parse top_10 string (semicolon or JSON format) into {value: count} dict <<< #
def _parse_top10(s):
    if not s or s == '[]':
        return {}
    s = str(s).strip()
    if s.startswith('['):
        try:
            entries = json.loads(s)
            return {str(e.get('value', '')): int(e.get('count', 0)) for e in entries}
        except (json.JSONDecodeError, TypeError):
            return {}
    result = {}
    for entry in s.split('; '):
        entry = entry.strip()
        if '(' in entry and entry.endswith(')'):
            val = entry[:entry.rfind('(')].strip()
            try:
                cnt = int(entry[entry.rfind('(') + 1:-1].strip())
            except ValueError:
                continue
            result[val] = cnt
    return result

# >>> Compare two top_10 strings. Returns True if different <<< #
def _compare_top10(left_str, right_str):
    left = _parse_top10(left_str)
    right = _parse_top10(right_str)
    return left != right

# >>> Parse column mapping string like "AMT=amount,STATUS=status" into dict <<< #
def parse_col_map_string(col_map_str):
    if not col_map_str:
        return {}
    mappings = {}
    for pair in col_map_str.split(','):
        pair = pair.strip()
        if '=' not in pair:
            continue
        left, right = pair.split('=', 1)
        mappings[left.strip()] = right.strip()
    return mappings

# >>> Get column mapping between two tables <<< #
def get_column_mapping(db_path, table_left, table_right, pair_name=None, col_map_override=None):
    if col_map_override is not None:
        return col_map_override

    if pair_name:
        pair = get_table_pair(db_path, pair_name)
        if pair and pair.get('col_mappings'):
            return pair['col_mappings']

    pairs = list_table_pairs(db_path)
    for pair in pairs:
        if pair['table_left'] == table_left and pair['table_right'] == table_right:
            return pair.get('col_mappings', {})
        if pair['table_left'] == table_right and pair['table_right'] == table_left:
            mappings = pair.get('col_mappings', {})
            return {v: k for k, v in mappings.items()}

    return {}

# >>> Compare row counts between two tables <<< #
def compare_row_counts(db_path, table_left, table_right, from_date=None, to_date=None,
                       date_synonyms=None):
    rows_left = get_row_counts(db_path, table_left, from_date, to_date,
                               date_synonyms=date_synonyms)
    rows_right = get_row_counts(db_path, table_right, from_date, to_date,
                                date_synonyms=date_synonyms)

    dict_left = {dt: count for dt, count in rows_left if dt}
    dict_right = {dt: count for dt, count in rows_right if dt}
    all_dates = set(dict_left.keys()) | set(dict_right.keys())

    # >>> within the overlap window a single-side date is a mismatch with the other side = 0 <<< #
    ov_start = ov_end = None
    if dict_left and dict_right:
        ov_start = max(min(dict_left), min(dict_right))
        ov_end = min(max(dict_left), max(dict_right))
        if ov_start > ov_end:
            ov_start = ov_end = None

    def _in_overlap(dt):
        return ov_start is not None and ov_start <= dt <= ov_end

    only_left = []
    only_right = []
    matching = []
    mismatched = []

    for dt in sorted(all_dates):
        if dt in dict_left and dt in dict_right:
            if dict_left[dt] == dict_right[dt]:
                matching.append((dt, dict_left[dt]))
            else:
                mismatched.append((dt, dict_left[dt], dict_right[dt]))
        elif dt in dict_left:
            if _in_overlap(dt):
                mismatched.append((dt, dict_left[dt], 0))
            else:
                only_left.append((dt, dict_left[dt]))
        else:
            if _in_overlap(dt):
                mismatched.append((dt, 0, dict_right[dt]))
            else:
                only_right.append((dt, dict_right[dt]))

    summary = {
        "date_range_left": (
            min(dict_left.keys()) if dict_left else None,
            max(dict_left.keys()) if dict_left else None,
        ),
        "date_range_right": (
            min(dict_right.keys()) if dict_right else None,
            max(dict_right.keys()) if dict_right else None,
        ),
        "total_left": sum(dict_left.values()),
        "total_right": sum(dict_right.values()),
        "count_left": len(dict_left),
        "count_right": len(dict_right),
    }

    return {
        "only_left": only_left,
        "only_right": only_right,
        "matching": matching,
        "mismatched": mismatched,
        "summary": summary,
    }

# >>> Resolve column types to a normalized family (numeric/date/categorical) <<< #
def resolve_col_type(left_type, right_type, overrides=None, col_name=None):
    if overrides and col_name and col_name in overrides:
        return overrides[col_name]
    lf, rf = _col_family(left_type), _col_family(right_type)
    return lf if lf == rf else "categorical"

# >>> classify a numeric stat pair (mean/std): 'match' (exact / <=abs_tiny), 'warn'
#     (relative drift <= rel_warn — low-order digits, likely cross-platform float /
#     overflow, PASSES), 'diff' (real), None (not comparable). Relative tol is the
#     industrial way to compare aggregates whose magnitude varies by orders. <<< #
def _float_status(a, b, rel_warn=1e-6, abs_tiny=1e-9):
    if a is None or a == '' or b is None or b == '':
        return None
    try:
        lf, rf = float(a), float(b)
    except (ValueError, TypeError):
        return 'match' if str(a).strip() == str(b).strip() else 'diff'
    diff = abs(lf - rf)
    if diff <= abs_tiny:
        return 'match'
    return 'warn' if diff / max(abs(lf), abs(rf), 1.0) <= rel_warn else 'diff'

# >>> Compare column statistics between two tables <<< #
def compare_column_stats(
    db_path, table_left, table_right,
    columns=None, col_mappings=None,
    from_date=None, to_date=None,
    matched_dates=None,
    col_type_overrides=None,
    level="fine",
):
    stats_left = get_col_stats(
        db_path, table_left, from_date=from_date, to_date=to_date)
    stats_right = get_col_stats(
        db_path, table_right, from_date=from_date, to_date=to_date)

    def organize_stats(stats):
        organized = {}
        for stat in stats:
            col = stat['column_name']
            raw = stat['dt']
            if raw:
                try:
                    dt = parse_date(raw)
                except ValueError:
                    dt = raw
            else:
                dt = raw
            if col not in organized:
                organized[col] = {}
            organized[col][dt] = stat
        return organized

    left_by_col = organize_stats(stats_left)
    right_by_col = organize_stats(stats_right)

    if columns is None:
        columns = list(left_by_col.keys())
    if col_mappings is None:
        col_mappings = {}

    result = {}

    for left_col in columns:
        right_col = col_mappings.get(left_col, left_col)

        if left_col not in left_by_col or right_col not in right_by_col:
            continue

        left_dates = left_by_col[left_col]
        right_dates = right_by_col[right_col]

        comparisons = []
        all_dates = set(left_dates.keys()) | set(right_dates.keys())
        if matched_dates is not None:
            all_dates = {d for d in all_dates
                         if d in _DATE_SENTINELS or d in matched_dates}

        for dt in sorted(all_dates):
            if dt not in left_dates or dt not in right_dates:
                continue

            left_stat = left_dates[dt]
            right_stat = right_dates[dt]

            left_type = left_stat["col_type"]
            right_type = right_stat["col_type"]
            resolved_type = resolve_col_type(left_type, right_type, col_type_overrides, left_col)

            l_total = _safe_int(left_stat["n_total"])
            r_total = _safe_int(right_stat["n_total"])
            l_missing = _safe_int(left_stat["n_missing"])
            r_missing = _safe_int(right_stat["n_missing"])
            l_unique = _safe_int(left_stat["n_unique"])
            r_unique = _safe_int(right_stat["n_unique"])

            comparison = {
                "dt": dt,
                "vintage_label": left_stat.get("vintage_label") or dt,
                "col_type": resolved_type,
                "col_type_left": left_type,
                "col_type_right": right_type,
                "left_col": left_col,
                "right_col": right_col,
                "n_total_left": l_total,
                "n_total_right": r_total,
                "n_total_diff": r_total - l_total,
                "n_missing_left": l_missing,
                "n_missing_right": r_missing,
                "n_missing_diff": r_missing - l_missing,
                "n_unique_left": l_unique,
                "n_unique_right": r_unique,
                "n_unique_diff": r_unique - l_unique,
            }

            def _norm_min_max(v):
                if v is None:
                    return ""
                return str(v).strip()
            l_min_key = _stat_match_key(left_stat["min_val"], "mn")
            r_min_key = _stat_match_key(right_stat["min_val"], "mn")
            l_max_key = _stat_match_key(left_stat["max_val"], "mx")
            r_max_key = _stat_match_key(right_stat["max_val"], "mx")
            comparison.update({
                "min_left": _norm_min_max(left_stat["min_val"]),
                "min_right": _norm_min_max(right_stat["min_val"]),
                "min_match": (l_min_key == r_min_key) if (l_min_key is not None and r_min_key is not None) else None,
                "max_left": _norm_min_max(left_stat["max_val"]),
                "max_right": _norm_min_max(right_stat["max_val"]),
                "max_match": (l_max_key == r_max_key) if (l_max_key is not None and r_max_key is not None) else None,
            })

            # >>> detect date-shaped columns to skip length signals <<< #
            def _is_date_string(v):
                if v is None or v == '':
                    return False
                s = str(v).strip()
                if not any(c.isalpha() or c in '-/: ' for c in s):
                    return False
                try:
                    from .date_utils import parse_date
                    parse_date(s)
                    return True
                except (ValueError, TypeError):
                    return False
            is_date_col = (
                _is_date_string(left_stat.get("min_val")) and
                _is_date_string(right_stat.get("min_val"))
            ) or (
                _is_date_string(left_stat.get("max_val")) and
                _is_date_string(right_stat.get("max_val"))
            )

            for field in ("length_avg", "length_max", "hash_sum"):
                raw_l = left_stat.get(field)
                raw_r = right_stat.get(field)
                lv = raw_l.strip() if isinstance(raw_l, str) else (raw_l or "")
                rv = raw_r.strip() if isinstance(raw_r, str) else (raw_r or "")
                if field in ("length_avg", "length_max") and is_date_col:
                    comparison[f"{field}_left"] = lv
                    comparison[f"{field}_right"] = rv
                    comparison[f"{field}_match"] = None
                    continue
                if lv and rv:
                    if field == "hash_sum":
                        try:
                            lf, rf = float(lv), float(rv)
                            denom = max(abs(lf), abs(rf), 1.0)
                            match = abs(lf - rf) / denom < 1e-9
                        except (ValueError, TypeError):
                            match = (lv == rv)
                    else:
                        try:
                            match = abs(float(lv) - float(rv)) < 1e-4
                        except (ValueError, TypeError):
                            match = (lv == rv)
                elif lv or rv:
                    match = False
                else:
                    match = None
                comparison[f"{field}_left"] = lv
                comparison[f"{field}_right"] = rv
                comparison[f"{field}_match"] = match

            if level == "fine":
                m_status = _float_status(left_stat["mean"], right_stat["mean"])
                s_status = _float_status(left_stat["std"], right_stat["std"])
                comparison.update({
                    "mean_left": left_stat["mean"],
                    "mean_right": right_stat["mean"],
                    "mean_match": None if m_status is None else (m_status != "diff"),
                    "mean_warn": m_status == "warn",
                    "std_left": left_stat["std"],
                    "std_right": right_stat["std"],
                    "std_match": None if s_status is None else (s_status != "diff"),
                    "std_warn": s_status == "warn",
                })
                if resolved_type != "numeric":
                    comparison.update({
                        "top_10_left": left_stat.get("top_10", ""),
                        "top_10_right": right_stat.get("top_10", ""),
                    })

            comparisons.append(comparison)

        if comparisons:
            result[left_col] = comparisons

    return result

# >>> field-key → diff predicate; single source of truth across helpers <<< #
_DIFF_PREDICATES = {
    'n_total':    lambda c: c.get('n_total_diff', 0) != 0,
    'n_missing':  lambda c: c.get('n_missing_diff', 0) != 0,
    'n_unique':   lambda c: c.get('n_unique_diff', 0) != 0,
    'mean':       lambda c: c.get('mean_match') is False,
    'std':        lambda c: c.get('std_match') is False,
    'min':        lambda c: c.get('min_match') is False,
    'max':        lambda c: c.get('max_match') is False,
    'length_avg': lambda c: c.get('length_avg_match') is False,
    'length_max': lambda c: c.get('length_max_match') is False,
    'hash_sum':   lambda c: c.get('hash_sum_match') is False,
}

# >>> Check if a column comparison has any statistical differences <<< #
def _has_col_differences(comp, compare_fields=None):
    if compare_fields is None:
        keys = _DIFF_PREDICATES.keys()
    else:
        active = set(compare_fields)
        keys = [k for k in _DIFF_PREDICATES if k in active]
    for k in keys:
        if _DIFF_PREDICATES[k](comp):
            return True
    return False

# >>> Compare two column dictionaries and match by case-insensitive name <<< #
def match_columns_from_dicts(left_cols, right_cols, left_label="left", right_label="right", outfile=None):
    left_lower = {k.lower(): k for k in left_cols}
    right_lower = {k.lower(): k for k in right_cols}

    matched = {}
    left_only = []
    right_only = []

    for lower_name, left_name in sorted(left_lower.items()):
        if lower_name in right_lower:
            right_name = right_lower[lower_name]
            matched[left_name] = right_name
        else:
            left_only.append({"name": left_name, "type": left_cols[left_name]})

    for lower_name, right_name in sorted(right_lower.items()):
        if lower_name not in left_lower:
            right_only.append({"name": right_name, "type": right_cols[right_name]})

    result = {
        "matched": matched,
        f"{left_label}_only": left_only,
        f"{right_label}_only": right_only,
        "manual_mapping": {},
    }

    if outfile:
        with open(outfile, 'w', encoding='utf-8') as f:
            json.dump(result, f, indent=2)
        print(f"\nWritten to: {outfile}")

    return result

# >>> transform value via a wildcard rule: `*` capture and paired (a|b)->(A|B) alternation <<< #
def _wildcard_transform(value, pat_from, pat_to):
    has_star = '*' in pat_from
    has_alt = '(' in pat_from and '(' in pat_to
    if not has_star and not has_alt:
        return pat_to if value == pat_from else None

    if has_alt:
        lm = re.search(r'\(([^)]+)\)', pat_from)
        rm = re.search(r'\(([^)]+)\)', pat_to)
        if lm and rm:
            left_alts = lm.group(1).split('|')
            right_alts = rm.group(1).split('|')
            alt_re = '(' + '|'.join(re.escape(a) for a in left_alts) + ')'
            full = re.sub(r'\([^)]+\)', alt_re, pat_from, count=1).replace('*', '(.*)')
            m = re.fullmatch(full, value, re.IGNORECASE)
            if not m:
                return None
            captured = m.group(1)
            idx = next((i for i, a in enumerate(left_alts)
                        if a.lower() == captured.lower()), -1)
            if idx < 0 or idx >= len(right_alts):
                return None
            result = re.sub(r'\([^)]+\)', right_alts[idx], pat_to, count=1)
            if m.lastindex and m.lastindex >= 2 and m.group(2) is not None:
                result = result.replace('*', m.group(2))
            return result

    regex = ''.join('(.*)' if ch == '*' else re.escape(ch) for ch in pat_from)
    m = re.fullmatch(regex, value, re.IGNORECASE)
    if not m:
        return None
    return pat_to.replace('*', m.group(1))

# >>> accept JS-style $1 / $<name> backrefs in a regex replacement template <<< #
def _to_py_backrefs(template):
    template = re.sub(r'\$(\d+)', r'\\\1', template)
    template = re.sub(r'\$<([^>]+)>', r'\\g<\1>', template)
    return template

# >>> Evaluate wildcard/regex rules against unmatched columns <<< #
def apply_column_rules(rules, unmatched_left, unmatched_right):
    new_mappings = {}
    rule_sources = {}
    right_set = set(unmatched_right)

    for rule_idx, rule in enumerate(rules):
        pat_left = rule.get("pattern_left", "")
        pat_right = rule.get("pattern_right", "")
        rule_type = rule.get("type", "wildcard")

        for left_col in list(unmatched_left):
            if left_col in new_mappings:
                continue

            if rule_type == "wildcard":
                transformed = _wildcard_transform(left_col, pat_left, pat_right)
                if transformed and transformed in right_set:
                    new_mappings[left_col] = transformed
                    rule_sources[left_col] = f"rule:{rule_idx}"
                    right_set.discard(transformed)
            elif rule_type == "regex":
                try:
                    m = re.fullmatch(pat_left, left_col, re.IGNORECASE)
                    if m:
                        transformed = m.expand(_to_py_backrefs(pat_right))
                        if transformed in right_set:
                            new_mappings[left_col] = transformed
                            rule_sources[left_col] = f"rule:{rule_idx}"
                            right_set.discard(transformed)
                except re.error:
                    continue

    return new_mappings, rule_sources

_SUPPORTED = ("aws", "oracle", "databricks", "csv")

# >>> source -> SQL dialect string used by quote_ident / hash16 exprs <<< #
def _dialect(source):
    return {"aws": "athena", "csv": "duckdb", "oracle": "oracle",
            "databricks": "databricks", "hadoop": "hadoop"}.get((source or "").lower())

# >>> per-value 16-bit MD5 hash expression matching each platform's col extraction <<< #
def _hash16_expr(source, col_q):
    src = (source or "").lower()
    if src == "oracle":
        from .platforms.oracle import _oracle_portable_hash16
        return _oracle_portable_hash16(col_q)
    if src == "aws":
        from .platforms.athena import _athena_portable_hash16
        return _athena_portable_hash16(col_q)
    if src == "databricks":
        from .platforms.databricks import _databricks_portable_hash16
        return _databricks_portable_hash16(col_q)
    if src == "csv":
        return (f"CAST(('0x' || substr(md5(trim(CAST({col_q} AS VARCHAR))), "
                f"1, 4)) AS INTEGER)")
    if src == "hadoop":
        return (f"CAST(conv(substr(lower(md5(trim(CAST({col_q} AS STRING)))), "
                f"1, 4), 16, 10) AS BIGINT)")
    return "NULL"

HASH_BUCKET_SPACE = 2 ** 32

# >>> 32-bit MD5 bucket of an already-normalized string expr (uniform 0..2^32-1) for cross-backend fraction sampling <<< #
def _hash32_expr(dialect, str_expr):
    d = (dialect or "").lower()
    if d == "oracle":
        return (f"TO_NUMBER(SUBSTR(STANDARD_HASH("
                f"UTL_I18N.STRING_TO_RAW({str_expr}, 'AL32UTF8'), 'MD5'), 1, 8), "
                f"'XXXXXXXX')")
    if d in ("athena", "trino", "aws"):
        return f"from_base(SUBSTR(LOWER(TO_HEX(MD5(TO_UTF8({str_expr})))), 1, 8), 16)"
    if d in ("databricks", "spark", "hadoop", "hive"):
        return f"CAST(conv(SUBSTR(LOWER(md5({str_expr})), 1, 8), 16, 10) AS BIGINT)"
    if d == "duckdb":
        return f"CAST(('0x' || SUBSTR(md5({str_expr}), 1, 8)) AS BIGINT)"
    raise ValueError(f"no hash-bucket expr for dialect {dialect!r}")

# >>> Python reference for _hash32_expr: 32-bit MD5 bucket of a normalized string <<< #
def hash32_bucket(norm_value):
    import hashlib
    if norm_value is None:
        return None
    return int(hashlib.md5(str(norm_value).encode("utf-8")).hexdigest()[:8], 16)

# >>> (from_ref, cte_prefix) for a table, honoring processed CTEs where supported <<< #
def _full_table_and_cte(tbl):
    src = (tbl.get("source") or "").lower()
    if src == "databricks":
        from .platforms.databricks import _full_table
        return _full_table(tbl), ""
    if src == "csv":
        cp = tbl.get("conn_macro", "") or ""
        return resolve_table(tbl, f"'{cp}'" if cp else tbl.get("table", ""))
    if src in ("aws", "hadoop"):
        cm = tbl.get("conn_macro", "") or ""
        tn = tbl.get("table", "")
        return resolve_table(tbl, f"{cm}.{tn}" if cm else tn)
    return resolve_table(tbl)

# >>> run a SELECT against the table's backend; return list[dict] (lowercase cols) <<< #
def _run_query(tbl, sql):
    src = (tbl.get("source") or "").lower()
    cm = tbl.get("conn_macro", "") or ""
    if src == "aws":
        from .platforms.athena import _query_athena
        df = _query_athena(sql, cm)
        return df.to_dict("records")
    if src == "databricks":
        from .platforms.databricks import _query_databricks
        return _query_databricks(sql, cm)
    if src == "csv":
        from .platforms.csv_duckdb import _query_duckdb
        return _query_duckdb(sql, cm)
    if src == "hadoop":
        from .platforms.hadoop import _query_hadoop
        return _query_hadoop(sql, cm)
    if src == "oracle":
        from .db import oracle_connect
        conn = oracle_connect(cm)
        if conn is None:
            return []
        try:
            cur = conn.cursor()
            cur.execute(sql)
            cols = [d[0].lower() for d in cur.description]
            rows = [dict(zip(cols, r)) for r in cur.fetchall()]
            cur.close()
            return rows
        finally:
            conn.close()
    raise ValueError(f"sample-hash: unsupported source {src!r} "
                     f"(supported: {', '.join(_SUPPORTED)})")

def _limit(source, n):
    return f"FETCH FIRST {n} ROWS ONLY" if (source or "").lower() == "oracle" else f"LIMIT {n}"

def _lit(v):
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, (int, float)):
        return repr(v)
    return "'" + str(v).replace("'", "''") + "'"

# >>> sample up to n distinct key tuples from the left side <<< #
def _sample_keys(tbl, key_cols, n):
    src = tbl.get("source", "")
    d = _dialect(src)
    ft, cte = _full_table_and_cte(tbl)
    keys_q = ", ".join(quote_ident(k, d) for k in key_cols)
    notnull = " AND ".join(f"{quote_ident(k, d)} IS NOT NULL" for k in key_cols)
    base_where = (tbl.get("where") or "").strip() or "1=1"
    sql = (f"{cte}SELECT DISTINCT {keys_q} FROM {ft} "
           f"WHERE ({base_where}) AND {notnull} {_limit(src, n)}")
    rows = _run_query(tbl, sql)
    out = []
    for r in rows:
        out.append(tuple(r[k.lower()] for k in key_cols))
    return out

# >>> all distinct key tuples from one side (no limit), in the given key order <<< #
def _all_keys(tbl, key_cols):
    src = tbl.get("source", "")
    d = _dialect(src)
    ft, cte = _full_table_and_cte(tbl)
    keys_q = ", ".join(quote_ident(k, d) for k in key_cols)
    notnull = " AND ".join(f"{quote_ident(k, d)} IS NOT NULL" for k in key_cols)
    base_where = (tbl.get("where") or "").strip() or "1=1"
    sql = (f"{cte}SELECT DISTINCT {keys_q} FROM {ft} "
           f"WHERE ({base_where}) AND {notnull}")
    rows = _run_query(tbl, sql)
    return [tuple(r[k.lower()] for k in key_cols) for r in rows]

# >>> deterministic hash-spread sample of up to n distinct key tuples from the left (for col-sample) <<< #
def sample_keys_hashed(tbl, key_cols, n):
    src = tbl.get("source", "")
    d = _dialect(src)
    ft, cte = _full_table_and_cte(tbl)
    keys_q = [quote_ident(k, d) for k in key_cols]
    notnull = " AND ".join(f"{kq} IS NOT NULL" for kq in keys_q)
    base_where = (tbl.get("where") or "").strip() or "1=1"
    order_expr = _hash16_expr(src, keys_q[0])
    sql = (f"{cte}SELECT {', '.join(keys_q)} FROM {ft} "
           f"WHERE ({base_where}) AND {notnull} "
           f"GROUP BY {', '.join(keys_q)} "
           f"ORDER BY {order_expr} {_limit(src, n)}")
    rows = _run_query(tbl, sql)
    return [tuple(r[k.lower()] for k in key_cols) for r in rows]

def _key_in_clause(key_cols_q, key_tuples):
    if not key_tuples:
        return "1=0"
    if len(key_cols_q) == 1:
        vals = ", ".join(_lit(t[0]) for t in key_tuples)
        return f"{key_cols_q[0]} IN ({vals})"
    ors = []
    for t in key_tuples:
        ands = " AND ".join(f"{kq} = {_lit(v)}" for kq, v in zip(key_cols_q, t))
        ors.append(f"({ands})")
    return "(" + " OR ".join(ors) + ")"

# >>> fetch value + hash16 for each (key, column) on one side; keyed by key tuple <<< #
def _fetch_side(tbl, key_cols, value_cols, key_tuples):
    src = tbl.get("source", "")
    d = _dialect(src)
    ft, cte = _full_table_and_cte(tbl)
    key_cols_q = [quote_ident(k, d) for k in key_cols]
    parts = [f'{kq} AS "__k{i}"' for i, kq in enumerate(key_cols_q)]
    for j, c in enumerate(value_cols):
        cq = quote_ident(c, d)
        parts.append(f'{cq} AS "__v{j}"')
        parts.append(f'{_hash16_expr(src, cq)} AS "__h{j}"')
    where = _key_in_clause(key_cols_q, key_tuples)
    base_where = (tbl.get("where") or "").strip()
    if base_where:
        where = f"({base_where}) AND {where}"
    sql = f"{cte}SELECT {', '.join(parts)} FROM {ft} WHERE {where}"
    rows = _run_query(tbl, sql)
    by_key = {}
    for r in rows:
        k = tuple(r.get(f"__k{i}") for i in range(len(key_cols)))
        by_key[tuple(str(x) for x in k)] = r
    return by_key

# >>> case-insensitive set of columns recorded for a table in _column_meta (empty if none) <<< #
def _present_cols(project_db, tbl):
    try:
        meta = get_column_meta(project_db, qualified_name(tbl))
    except Exception:
        meta = []
    return {m["column_name"].lower() for m in meta if m.get("column_name")}

# >>> sample-hash one pair: returns (header, rows) for the Excel sheet <<< #
def _sample_pair(project_db, pair_name, pair_cfg, key_cols, n):
    left = dict(pair_cfg["left"])
    right = dict(pair_cfg["right"])
    for side in (left, right):
        side.setdefault("name", "")
    col_map = pair_cfg.get("col_map") or get_pair_col_map_from_db(project_db, pair_name) or {}
    if not col_map:
        print(f"  {pair_name}: no col_map; run `dtrack match-columns` first. Skipping.")
        return None
    for src in (left.get("source", ""), right.get("source", "")):
        if (src or "").lower() not in _SUPPORTED:
            print(f"  {pair_name}: source {src!r} not supported by sample-hash "
                  f"({', '.join(_SUPPORTED)}). Skipping.")
            return None

    left_present = _present_cols(project_db, left)
    right_present = _present_cols(project_db, right)

    def _ok(col, present):
        return (not present) or (str(col).lower() in present)

    left_keys = key_cols
    right_keys = [col_map.get(k, k) for k in key_cols]
    value_pairs, dropped = [], []
    for l, r in col_map.items():
        if l in key_cols:
            continue
        if _ok(l, left_present) and _ok(r, right_present):
            value_pairs.append((l, r))
        else:
            dropped.append(l if l == r else f"{l}->{r}")
    if dropped:
        print(f"  {pair_name}: skipping {len(dropped)} column(s) absent on one "
              f"side: {', '.join(dropped)}")
    if not value_pairs:
        print(f"  {pair_name}: no matched columns present on both sides. Skipping.")
        return None
    left_vals = [l for l, _ in value_pairs]
    right_vals = [r for _, r in value_pairs]

    key_tuples = _sample_keys(left, left_keys, n)
    if not key_tuples:
        print(f"  {pair_name}: no key rows sampled from left. Skipping.")
        return None
    norm_keys = [tuple(str(x) for x in t) for t in key_tuples]

    left_by = _fetch_side(left, left_keys, left_vals, key_tuples)
    right_by = _fetch_side(right, right_keys, right_vals, key_tuples)

    cols_hdr = ["key"]
    for l, r in value_pairs:
        label = l if l == r else f"{l}->{r}"
        cols_hdr += [f"{label}.L", f"{label}.Lh16", f"{label}.R", f"{label}.Rh16", f"{label}.match"]

    rows = []
    for nk in norm_keys:
        lrow = left_by.get(nk, {})
        rrow = right_by.get(nk, {})
        out = ["|".join(nk)]
        for j in range(len(value_pairs)):
            lv = lrow.get(f"__v{j}")
            rv = rrow.get(f"__v{j}")
            lh = lrow.get(f"__h{j}")
            rh = rrow.get(f"__h{j}")
            match = (lh is not None and rh is not None and str(lh) == str(rh))
            out += [_disp(lv), _disp(lh), _disp(rv), _disp(rh), "OK" if match else "DIFF"]
        rows.append(out)
    return cols_hdr, rows

def _disp(v):
    return "" if v is None else str(v)

# >>> n=-1: full key-set anti-join — key tuples present on one side but not the other <<< #
def _reconcile_keys_pair(project_db, pair_name, pair_cfg, key_cols):
    left = dict(pair_cfg["left"])
    right = dict(pair_cfg["right"])
    for side in (left, right):
        side.setdefault("name", "")
    for src in (left.get("source", ""), right.get("source", "")):
        if (src or "").lower() not in _SUPPORTED:
            print(f"  {pair_name}: source {src!r} not supported by sample-hash "
                  f"({', '.join(_SUPPORTED)}). Skipping.")
            return None
    col_map = pair_cfg.get("col_map") or get_pair_col_map_from_db(project_db, pair_name) or {}
    left_keys = key_cols
    right_keys = [col_map.get(k, k) for k in key_cols]
    left_set = {tuple(str(x) for x in t) for t in _all_keys(left, left_keys)}
    right_set = {tuple(str(x) for x in t) for t in _all_keys(right, right_keys)}
    left_only = sorted(left_set - right_set)
    right_only = sorted(right_set - left_set)
    print(f"  {pair_name}: left={len(left_set)} right={len(right_set)} distinct keys; "
          f"{len(left_only)} left-only, {len(right_only)} right-only")
    header = list(key_cols) + ["side"]
    rows = ([list(t) + ["LEFT_ONLY"] for t in left_only]
            + [list(t) + ["RIGHT_ONLY"] for t in right_only])
    return header, rows

# >>> write one sheet per pair; highlight DIFF rows; bold header <<< #
def _write_excel(out_path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    diff_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    bad_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
    bold = Font(bold=True)
    for name, (header, rows) in sheets.items():
        ws = wb.create_sheet(title=name[:31])
        ws.append(header)
        for c in ws[1]:
            c.font = bold
        match_cols = [i for i, h in enumerate(header) if h.endswith(".match")]
        for row in rows:
            ws.append(row)
            r = ws[ws.max_row]
            for mc in match_cols:
                if row[mc] == "DIFF":
                    for c in r:
                        c.fill = diff_fill
                    r[mc].fill = bad_fill
    wb.save(out_path)
    return out_path

# >>> CLI entry: sample N keys/pair, dump per-column value+backend hash16 to Excel <<< #
def sample_hash(project_db, config_path, key, pair=None, n=20, out=None):
    if not key:
        raise ValueError("sample-hash requires --key <col>[,<col2>...]")
    key_cols = [k.strip() for k in str(key).split(",") if k.strip()]
    config = load_unified_config(config_path)
    pairs = config.get("pairs", {})
    if pair:
        if pair not in pairs:
            raise ValueError(f"pair {pair!r} not in config")
        items = [(pair, pairs[pair])]
    else:
        items = [(p, c) for p, c in pairs.items() if not c.get("skip")]

    reconcile = (n == -1)
    sheets = {}
    for pname, pcfg in items:
        mode = "all rows, key-reconcile" if reconcile else f"n={n}"
        print(f"\n--- sample-hash: {pname} (key={','.join(key_cols)}, {mode}) ---")
        try:
            if reconcile:
                res = _reconcile_keys_pair(project_db, pname, pcfg, key_cols)
            else:
                res = _sample_pair(project_db, pname, pcfg, key_cols, n)
        except Exception as e:
            print(f"  {pname}: failed: {e}")
            continue
        if res:
            sheets[pname] = res
            if not reconcile:
                diffs = sum(1 for row in res[1] for v in row if v == "DIFF")
                print(f"  {pname}: {len(res[1])} rows sampled, {diffs} cell hash diffs")

    if not sheets:
        print("\nNo pairs produced output (need col_map + supported sources).")
        return None
    out = out or "sample_hash.xlsx"
    _write_excel(out, sheets)
    print(f"\nExcel report: {out}")
    return out
